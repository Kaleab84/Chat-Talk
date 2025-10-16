"""Exercise the DocumentProcessor image extraction pipeline on a source DOC file.

This mirrors the exploratory `tmp_analyze.py` helper, but wrapped with a CLI so
we can quickly verify how the processor surfaces inline images and the storage
placeholders that will eventually be persisted.

Usage:
    python -m scripts.doc_debug.analyze_processor_images --doc path/to/file.DOC
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

# Ensure project imports resolve when executed directly.
REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.append(str(REPO_ROOT))

from app.services.document_processor import DocumentProcessor
from docx import Document


def analyze_images(doc_path: Path, doc_slug: str = "debug") -> None:
    """Convert a DOC file and display the image placeholders picked up."""
    processor = DocumentProcessor()
    converted_docx = processor._convert_doc_to_docx(doc_path)
    document = Document(converted_docx)

    ignore_rel_ids = processor._collect_header_footer_image_rel_ids(document)
    images, rel_to_meta = processor._extract_images(
        document,
        ignore_rel_ids,
        doc_slug=doc_slug,
    )

    print(f"Found {len(images)} extracted images.")
    print(f"Relationship map size: {len(rel_to_meta)}")

    for idx, paragraph in enumerate(document.paragraphs):
        paths = processor._inline_images_from_paragraph(paragraph, rel_to_meta)
        if paths:
            print(f"Paragraph {idx} references inline image(s):")
            for path in paths:
                print(f"  - {path}")

    print("Analysis complete.")


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Run DocumentProcessor image extraction on a DOC/DOCX file."
    )
    parser.add_argument(
        "--doc",
        type=Path,
        required=True,
        help="Path to the legacy document to analyze.",
    )
    parser.add_argument(
        "--slug",
        type=str,
        default="debug",
        help="Optional slug to pass through to the processor.",
    )
    return parser


def main() -> None:
    parser = _build_parser()
    args = parser.parse_args()
    analyze_images(args.doc.expanduser(), args.slug)


if __name__ == "__main__":
    main()
