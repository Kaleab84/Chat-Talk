# app/services/document_processor.py
from __future__ import annotations
#TODO: Add an engpoint to upload a file to be ingested. Like a drag and drop featue or selct from folder

import hashlib
import logging
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from docx import Document
from docx.document import Document as _DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.config import settings  # if you need settings in the future

# Ensure python-docx exposes VML namespace during xpath calls (needed for legacy images)
try:
    from docx.oxml.ns import nsmap as _docx_nsmap

    if "v" not in _docx_nsmap:
        _docx_nsmap["v"] = "urn:schemas-microsoft-com:vml"
    if "o" not in _docx_nsmap:
        _docx_nsmap["o"] = "urn:schemas-microsoft-com:office:office"
except Exception:
    # Namespace patching is best-effort; absence just means legacy images may be skipped.
    pass

logger = logging.getLogger(__name__)

# Optional Windows COM for high-fidelity DOC->DOCX
try:
    import win32com.client as win32  # type: ignore
    _HAS_WIN32 = True
except (ImportError, ModuleNotFoundError):
    _HAS_WIN32 = False


def _ensure_word_application():
    """Return a running Word COM instance, rebuilding the gencache if necessary."""
    if not _HAS_WIN32:
        raise RuntimeError("win32com is not available on this system.")

    try:
        return win32.gencache.EnsureDispatch("Word.Application")  # type: ignore[name-defined]
    except AttributeError as exc:
        # Known pywin32 issue: corrupted generated cache lacking CLSIDToClassMap.
        if "CLSIDToClassMap" not in str(exc):
            raise
        logger.warning("win32com cache appears corrupt; attempting to rebuild.")
        try:
            from win32com.client import gencache  # type: ignore

            # Allow rebuild to overwrite cached modules.
            if hasattr(gencache, "is_readonly"):
                gencache.is_readonly = False  # type: ignore[attr-defined]
            gencache.Rebuild()
        except Exception as rebuild_exc:
            raise RuntimeError(
                "Failed to rebuild the Word COM type library cache. "
                "Try deleting the %LOCALAPPDATA%\\Temp\\gen_py directory manually."
            ) from rebuild_exc
        # Second attempt after rebuild
        return win32.gencache.EnsureDispatch("Word.Application")  # type: ignore[name-defined]


# -------------------------
# Utilities
# -------------------------

def _is_valid_docx(path: Path) -> bool:
    """Heuristic: is file a valid OOXML docx (zip with 'word/' entries)?"""
    try:
        with zipfile.ZipFile(path) as z:
            return any(n.startswith("word/") for n in z.namelist())
    except Exception:
        return False


def _norm_text(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip())


def _hash_bytes(b: bytes) -> str:
    return hashlib.sha1(b).hexdigest()


def _slugify(name: str) -> str:
    name = name.strip().lower()
    name = re.sub(r"[^\w\-]+", "-", name)
    name = re.sub(r"-{2,}", "-", name).strip("-")
    return name or "document"


# -------------------------
# Image record
# -------------------------

@dataclass
class _ImageRecord:
    image_id: str
    extension: str
    data: bytes
    source_rel_id: Optional[str] = None
    width_px: Optional[int] = None
    height_px: Optional[int] = None
    # Friendly naming hints (filled later)
    doc_slug: Optional[str] = None
    seq: Optional[int] = None
    suggested_name: Optional[str] = None

    @property
    def placeholder_path(self) -> str:
        return f"images/{self.image_id}{self.extension}"


# -------------------------
# Processor
# -------------------------

class DocumentProcessor:
    """
    Returns dict with keys your pipeline expects:
      {
        'success': True/False,
        'doc_id': <doc_slug>,
        'source': <input path>,
        'doc_slug': <safe basename>,                # extra hint for naming
        'sections': [ { section_id, title, level, blocks, suggested_name? ... }, ... ],
        'images':   [ { image_id, extension, data, suggested_name?, ... }, ... ],
        'chunks':   [ { chunk_id, section_id, text, image_paths: [...] }, ... ],
      }
    """

    # -------- Public entry --------

    def process_document(self, file_path: Path) -> Dict[str, Any]:
        try:
            if not file_path.exists():
                return {"success": False, "error": f"File not found: {file_path}", "source": str(file_path)}

            ext = file_path.suffix.lower()
            doc_slug = _slugify(file_path.stem)
            doc_id = doc_slug or str(uuid.uuid4())

            # Normalize: convert if .doc OR a bogus .docx (renamed .doc)
            if ext == ".doc" or (ext == ".docx" and not _is_valid_docx(file_path)):
                tmp_docx = self._convert_doc_to_docx(file_path)
                try:
                    return self._process_docx(tmp_docx, source=file_path, doc_id=doc_id)
                finally:
                    # cleanup temp conversion directory if we made one
                    try:
                        tmp_dir = tmp_docx.parent
                        if tmp_docx.exists():
                            tmp_docx.unlink(missing_ok=True)
                        if str(tmp_dir).endswith("__doc_conv"):
                            shutil.rmtree(tmp_dir, ignore_errors=True)
                    except Exception:
                        pass

            if ext == ".docx":
                return self._process_docx(file_path, source=file_path, doc_id=doc_id)

            return {"success": False, "error": f"Unsupported file type: {ext}", "source": str(file_path)}

        except Exception as e:
            logger.exception("process_document failed")
            return {"success": False, "error": str(e), "source": str(file_path)}

    # -------- Conversion --------

    def _convert_doc_to_docx(self, doc_path: Path) -> Path:
        """
        Prefer Word COM on Windows to embed linked images, fallback to LibreOffice headless.
        Returns a temporary DOCX path in a temp dir that ends with __doc_conv.
        """
        if os.name == "nt" and _HAS_WIN32:
            logger.info("Converting via Word COM: %s", doc_path)
            try:
                return self._convert_with_word_com(doc_path)
            except Exception as exc:
                logger.warning(
                    "Word COM conversion failed (%s); attempting LibreOffice fallback.", exc
                )

        logger.info("Converting via LibreOffice: %s", doc_path)
        return self._convert_with_libreoffice(doc_path)

    def _convert_with_word_com(self, doc_path: Path) -> Path:
        tmp_dir = Path(tempfile.mkdtemp(suffix="__doc_conv"))
        out_path = tmp_dir / f"{doc_path.stem}.docx"

        word = _ensure_word_application()
        word.Visible = False
        try:
            doc = word.Documents.Open(str(doc_path))
            # Update fields; force image embedding + break links so images are saved with document
            try:
                doc.Fields.Update()
            except Exception:
                pass
            try:
                count = doc.InlineShapes.Count
            except Exception:
                count = 0
            for i in range(1, count + 1):
                try:
                    ish = doc.InlineShapes.Item(i)
                    if hasattr(ish, "LinkFormat") and ish.LinkFormat is not None:
                        try:
                            ish.LinkFormat.SavePictureWithDocument = True
                        except Exception:
                            pass
                        try:
                            ish.LinkFormat.BreakLink()
                        except Exception:
                            pass
                except Exception:
                    pass

            # 16 = wdFormatXMLDocument
            doc.SaveAs(str(out_path), FileFormat=16)
            doc.Close(False)
            return out_path
        finally:
            word.Quit()

    def _convert_with_libreoffice(self, doc_path: Path) -> Path:
        tmp_dir = Path(tempfile.mkdtemp(suffix="__doc_conv"))
        try:
            cmd = [
                "soffice",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(tmp_dir),
                str(doc_path),
            ]
            subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            candidate = tmp_dir / f"{doc_path.stem}.docx"
            if not candidate.exists():
                raise RuntimeError(f"LibreOffice conversion failed to produce: {candidate}")
            return candidate
        except Exception:
            shutil.rmtree(tmp_dir, ignore_errors=True)
            raise

    # -------- DOCX pipeline --------

    def _process_docx(self, docx_path: Path, *, source: Path, doc_id: Optional[str] = None) -> Dict[str, Any]:
        if not _is_valid_docx(docx_path):
            raise RuntimeError(f"Invalid DOCX container: {docx_path}")

        doc = Document(str(docx_path))
        doc_slug = _slugify(Path(source).stem)
        doc_id = doc_id or doc_slug

        # Gather header/footer image relIds to ignore
        hf_rel_ids = self._collect_header_footer_image_rel_ids(doc)

        # Extract images (skip header/footer). Map relId -> (image_id, ext)
        images, rel_to_imgmeta = self._extract_images(doc, hf_rel_ids, doc_slug=doc_slug)

        # Build sections/blocks; includes inline image blocks with extension-correct placeholders
        sections = self._build_sections(doc, rel_to_imgmeta, doc_slug=doc_slug)

        # Build chunks from text/table blocks; attach image_paths
        chunks = self._build_chunks(sections)

        return {
            "success": True,
            "doc_id": doc_id,
            "doc_slug": doc_slug,  # naming hint for repository
            "source": str(source),
            "sections": sections,
            "images": [self._image_to_dict(img) for img in images],
            "chunks": chunks,
        }

    # -------- Header/Footer skip --------

    def _collect_header_footer_image_rel_ids(self, doc: _DocxDocument) -> set:
        """
        Scan all headers/footers for <a:blip r:embed="rId#"> and return rIds to skip.
        """
        ns = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        rel_ids = set()
        for sec in doc.sections:
            # Header
            try:
                hdr_part = sec.header.part
                for blip in hdr_part.element.xpath(".//a:blip", namespaces=ns):
                    rId = blip.get(f"{{{ns['r']}}}embed")
                    if rId:
                        rel_ids.add(rId)
            except Exception:
                pass
            # Footer
            try:
                ftr_part = sec.footer.part
                for blip in ftr_part.element.xpath(".//a:blip", namespaces=ns):
                    rId = blip.get(f"{{{ns['r']}}}embed")
                    if rId:
                        rel_ids.add(rId)
            except Exception:
                pass

        logger.info("Header/Footer images ignored (relIds): %d", len(rel_ids))
        return rel_ids

    # -------- Images --------

    def _extract_images(
        self, doc: _DocxDocument, ignore_rel_ids: set, *, doc_slug: str
    ) -> Tuple[List[_ImageRecord], Dict[str, Tuple[str, str]]]:
        """
        Returns:
          - list of _ImageRecord (dedup by bytes)
          - mapping: relId -> (image_id, extension) for inline references
        """
        related = doc.part.related_parts  # dict: rId -> Part
        rel_to_imgmeta: Dict[str, Tuple[str, str]] = {}
        images: List[_ImageRecord] = []
        seen_hashes = set()
        seq = 0

        for rId, part in related.items():
            if rId in ignore_rel_ids:
                continue
            ct = getattr(part, "content_type", "")
            if not ct.startswith("image/"):
                continue
            try:
                blob = part.blob
            except Exception:
                continue

            h = _hash_bytes(blob)
            if h in seen_hashes:
                continue
            seen_hashes.add(h)

            ext = self._ext_from_content_type(ct)
            img_id = h[:16]
            seq += 1

            rec = _ImageRecord(
                image_id=img_id,
                extension=ext,
                data=blob,
                source_rel_id=rId,
                doc_slug=doc_slug,
                seq=seq,
                suggested_name=f"{doc_slug}__img{seq:03d}{ext}",
            )
            images.append(rec)
            rel_to_imgmeta[rId] = (img_id, ext)

        logger.info("Extracted %d body images (after skipping header/footer)", len(images))
        return images, rel_to_imgmeta

    def _ext_from_content_type(self, ct: str) -> str:
        m = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/jpg": ".jpg",
            "image/gif": ".gif",
            "image/bmp": ".bmp",
            "image/tiff": ".tif",
            "image/x-wmf": ".wmf",
            "image/x-emf": ".emf",
        }
        return m.get(ct.lower(), ".bin")

    def _image_to_dict(self, img: _ImageRecord) -> Dict[str, Any]:
        return {
            "image_id": img.image_id,
            "extension": img.extension,
            "data": img.data,  # ingest step removes this after writing
            "width_px": img.width_px,
            "height_px": img.height_px,
            "doc_slug": img.doc_slug,
            "suggested_name": img.suggested_name,
            "seq": img.seq,
        }

    # -------- Sections & blocks --------

    def _build_sections(
        self,
        doc: _DocxDocument,
        rel_to_imgmeta: Dict[str, Tuple[str, str]],
        *,
        doc_slug: str,
    ) -> List[Dict[str, Any]]:
        """
        Construct sections using Heading styles; include:
          - text blocks
          - image blocks (inline & anchored) with correct extension placeholders
          - tables (as rows; images inside cells also captured)
        """
        # Flatten document body into items
        items: List[Tuple[str, Any]] = []
        body_elm = doc.element.body
        for child in body_elm.iterchildren():
            tag = child.tag.lower()
            if tag.endswith("p"):
                items.append(("p", Paragraph(child, doc)))
            elif tag.endswith("tbl"):
                items.append(("tbl", Table(child, doc)))

        sections: List[Dict[str, Any]] = []
        current: Optional[Dict[str, Any]] = None
        default_title = "Untitled"
        sec_idx = 0

        for kind, obj in items:
            if kind == "p":
                p: Paragraph = obj
                style_name = (p.style.name if p.style else "") or ""
                is_heading = bool(re.match(r"heading\s*\d+\Z", style_name.strip(), flags=re.I))

                if is_heading:
                    level_match = re.findall(r"\d+", style_name)
                    level = int(level_match[0]) if level_match else 1
                    title = _norm_text(p.text) or default_title
                    sec_idx += 1
                    current = {
                        "section_id": str(uuid.uuid4()),
                        "title": title,
                        "level": level,
                        "blocks": [],
                        "suggested_name": f"{doc_slug}__sec{sec_idx:03d}_{_slugify(title)[:40]}.json",
                        "doc_slug": doc_slug,
                        "index": sec_idx,
                    }
                    sections.append(current)
                    continue

                if current is None:
                    sec_idx += 1
                    current = {
                        "section_id": str(uuid.uuid4()),
                        "title": default_title,
                        "level": 1,
                        "blocks": [],
                        "suggested_name": f"{doc_slug}__sec{sec_idx:03d}_{_slugify(default_title)}.json",
                        "doc_slug": doc_slug,
                        "index": sec_idx,
                    }
                    sections.append(current)

                text = _norm_text(p.text)
                if text:
                    current["blocks"].append({"type": "text", "text": text})

                # Inline & anchored images in this paragraph (modern + VML)
                for path in self._inline_images_from_paragraph(p, rel_to_imgmeta):
                    current["blocks"].append({"type": "image", "path": path})

            else:  # table
                if current is None:
                    sec_idx += 1
                    current = {
                        "section_id": str(uuid.uuid4()),
                        "title": default_title,
                        "level": 1,
                        "blocks": [],
                        "suggested_name": f"{doc_slug}__sec{sec_idx:03d}_{_slugify(default_title)}.json",
                        "doc_slug": doc_slug,
                        "index": sec_idx,
                    }
                    sections.append(current)

                table: Table = obj
                table_rows: List[List[str]] = []
                for row in table.rows:
                    row_cells: List[str] = []
                    for cell in row.cells:
                        row_cells.append(_norm_text(cell.text))
                    table_rows.append(row_cells)

                current["blocks"].append({"type": "table", "rows": table_rows})

                # Also capture images inside table cells
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for path in self._inline_images_from_paragraph(p, rel_to_imgmeta):
                                current["blocks"].append({"type": "image", "path": path})

        if not sections:
            sec_idx = 1
            sections.append({
                "section_id": str(uuid.uuid4()),
                "title": default_title,
                "level": 1,
                "blocks": [],
                "suggested_name": f"{doc_slug}__sec{sec_idx:03d}_{_slugify(default_title)}.json",
                "doc_slug": doc_slug,
                "index": sec_idx,
            })

        return sections

    def _inline_images_from_paragraph(
        self,
        p: Paragraph,
        rel_to_imgmeta: Dict[str, Tuple[str, str]],
    ) -> List[str]:
        """
        Find images referenced in a paragraph via modern DrawingML (<a:blip r:embed>)
        and legacy VML (<v:imagedata r:id>), return placeholders with extension:
            images/<image_id>.<ext>
        """
        paths: List[str] = []

        def _xpath(element, query: str):
            """Call element.xpath with fallback for python-docx namespace quirks."""
            try:
                return element.xpath(query, namespaces={
                    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                    "v": "urn:schemas-microsoft-com:vml",
                })
            except TypeError:
                # python-docx>=1.1 overrides xpath(nsmap only); rely on patched ns map.
                return element.xpath(query)

        # Modern drawings
        try:
            for blip in _xpath(p._p, ".//a:blip"):
                rId = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if not rId:
                    continue
                meta = rel_to_imgmeta.get(rId)
                if meta:
                    img_id, ext = meta
                    paths.append(f"images/{img_id}{ext}")
        except Exception:
            pass

        # Legacy VML images
        try:
            for im in _xpath(p._p, ".//v:imagedata"):
                rId = im.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if not rId:
                    continue
                meta = rel_to_imgmeta.get(rId)
                if meta:
                    img_id, ext = meta
                    paths.append(f"images/{img_id}{ext}")
        except Exception:
            pass

        return paths

    # -------- Chunking --------

    def _build_chunks(self, sections: List[Dict[str, Any]], *, max_chars: int = 1200) -> List[Dict[str, Any]]:
        """
        Greedy character-based chunking over text and small tables.
        Attach image_paths encountered within the same section in order.
        """
        chunks: List[Dict[str, Any]] = []

        for sec in sections:
            sec_id = sec["section_id"]
            pending_text: List[str] = []
            pending_imgs: List[str] = []

            def _flush():
                if not pending_text and not pending_imgs:
                    return
                text = _norm_text(" ".join(pending_text))
                chunks.append({
                    "chunk_id": str(uuid.uuid4()),
                    "section_id": sec_id,
                    "text": text,
                    "image_paths": list(pending_imgs),
                })
                pending_text.clear()
                pending_imgs.clear()

            for blk in sec.get("blocks", []):
                btype = blk.get("type")
                if btype == "text":
                    t = _norm_text(blk.get("text", ""))
                    if not t:
                        continue
                    if sum(len(x) for x in pending_text) + len(t) + 1 > max_chars:
                        _flush()
                    pending_text.append(t)

                elif btype == "image":
                    path = blk.get("path")
                    if path:
                        pending_imgs.append(path)
                    # If text is already large, flush to avoid overlong chunks
                    if sum(len(x) for x in pending_text) > int(max_chars * 0.8):
                        _flush()

                elif btype == "table":
                    rows: List[List[str]] = blk.get("rows", [])
                    table_as_text = self._table_to_text(rows)
                    if table_as_text and len(table_as_text) < max_chars // 2:
                        if sum(len(x) for x in pending_text) + len(table_as_text) + 1 > max_chars:
                            _flush()
                        pending_text.append(table_as_text)
                    else:
                        _flush()
                        chunks.append({
                            "chunk_id": str(uuid.uuid4()),
                            "section_id": sec_id,
                            "text": (table_as_text or "")[:max_chars],
                            "image_paths": [],
                        })

            _flush()

        return chunks

    def _table_to_text(self, rows: List[List[str]]) -> str:
        if not rows:
            return ""
        return "\n".join(" | ".join((c or "").strip() for c in r) for r in rows)


# -------------------------
# Public module API (singleton)
# -------------------------

_default_processor = DocumentProcessor()

def process_document(file_path: Path) -> Dict[str, Any]:
    return _default_processor.process_document(file_path)



